"""
Moteur d'extraction FTTH APD : KMZ -> Excel (BPE_PBO, Support, Cable (Ligne), Synoptique).

Dépendances volontairement limitées pour rester compatible avec l'environnement
Python embarqué de QGIS : lxml, pandas, numpy, openpyxl (pas de scipy, pas de
beautifulsoup4).
"""
import zipfile
import os
import re
import math
import numpy as np
import pandas as pd
from lxml import etree
import lxml.html as LH
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Border, Side, Alignment, PatternFill
from openpyxl.utils import get_column_letter

NS = {"kml": "http://www.opengis.net/kml/2.2"}

# ============================================================
# 1) Lecture KMZ / parsing KML
# ============================================================

def load_kml_from_kmz(kmz_path):
    with zipfile.ZipFile(kmz_path, "r") as z:
        kml_names = [n for n in z.namelist() if n.lower().endswith(".kml")]
        if not kml_names:
            raise ValueError("Aucun fichier .kml trouvé dans le KMZ")
        main = next((n for n in kml_names if n.lower().endswith("doc.kml")), kml_names[0])
        return z.read(main)


def parse_description_table(description_html):
    """Parse une description au format tableau HTML 2 colonnes (clé/valeur)."""
    if not description_html:
        return {}
    try:
        doc = LH.fromstring(description_html)
    except Exception:
        return {}
    table = doc if doc.tag == "table" else doc.find(".//table")
    if table is None:
        return {}
    attrs = {}
    for tr in table.findall(".//tr"):
        cells = tr.findall("./td")
        if not cells:
            cells = tr.findall("./th")
        if len(cells) >= 2:
            key = "".join(cells[0].itertext()).strip()
            key = key.rstrip(":").strip()
            val = "".join(cells[1].itertext()).strip()
            if key:
                attrs[key] = val
    return attrs


def _direct_folder_name(pm):
    for ancestor in pm.iterancestors():
        if etree.QName(ancestor).localname == "Folder":
            name_el = ancestor.find("kml:name", NS)
            return name_el.text.strip() if name_el is not None and name_el.text else ""
    return ""


def parse_placemarks(kml_bytes):
    root = etree.fromstring(kml_bytes)
    results = []

    for pm in root.iter("{http://www.opengis.net/kml/2.2}Placemark"):
        name_el = pm.find("kml:name", NS)
        name = name_el.text.strip() if name_el is not None and name_el.text else ""

        desc_el = pm.find("kml:description", NS)
        description = desc_el.text if desc_el is not None else ""

        geom_type, coords = None, None
        point_el = pm.find(".//kml:Point/kml:coordinates", NS)
        line_el = pm.find(".//kml:LineString/kml:coordinates", NS)

        if point_el is not None and point_el.text:
            geom_type = "Point"
            coords = point_el.text.strip()
        elif line_el is not None and line_el.text:
            geom_type = "LineString"
            coords = line_el.text.strip()

        entry = {
            "name": name,
            "folder": _direct_folder_name(pm),
            "geometry_type": geom_type,
            "coordinates": coords,
        }
        entry.update(parse_description_table(description))

        for simple_data in pm.findall(".//kml:ExtendedData/kml:SchemaData/kml:SimpleData", NS):
            attr_name = simple_data.get("name")
            if attr_name:
                entry[attr_name] = simple_data.text.strip() if simple_data.text else ""

        results.append(entry)

    return results


def build_layers(kmz_path):
    kml_bytes = load_kml_from_kmz(kmz_path)
    data = parse_placemarks(kml_bytes)
    points = [d for d in data if d.get("geometry_type") == "Point"]
    lines = [d for d in data if d.get("geometry_type") == "LineString"]
    return pd.DataFrame(points), pd.DataFrame(lines)


# ============================================================
# 2) Géométrie : distances, projection, recherche spatiale (sans scipy)
# ============================================================

def _parse_lonlat(coord_str):
    if not coord_str:
        return None, None
    first = coord_str.strip().split()[0]
    parts = first.split(",")
    return float(parts[0]), float(parts[1])


def haversine_m(lon1, lat1, lon2, lat2):
    R = 6371000.0
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dphi = math.radians(lat2 - lat1)
    dlambda = math.radians(lon2 - lon1)
    a = math.sin(dphi / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dlambda / 2) ** 2
    return 2 * R * math.asin(math.sqrt(a))


def _coords_close(c1, c2, tolerance_m):
    return haversine_m(c1[0], c1[1], c2[0], c2[1]) <= tolerance_m


class SpatialIndex:
    """Index spatial brute-force (numpy) : suffisant pour quelques milliers de
    points, évite la dépendance à scipy.spatial.cKDTree."""

    def __init__(self, lons, lats, ref_lat=None):
        self.lons = np.asarray(lons, dtype=float)
        self.lats = np.asarray(lats, dtype=float)
        if ref_lat is None:
            ref_lat = float(np.mean(self.lats)) if len(self.lats) else 0.0
        self.m_per_deg_lat = 111320.0
        self.m_per_deg_lon = 111320.0 * math.cos(math.radians(ref_lat))
        self.x = self.lons * self.m_per_deg_lon
        self.y = self.lats * self.m_per_deg_lat

    def query_ball_point(self, lon, lat, r):
        if len(self.x) == 0:
            return []
        x = lon * self.m_per_deg_lon
        y = lat * self.m_per_deg_lat
        d = np.hypot(self.x - x, self.y - y)
        return np.where(d <= r)[0].tolist()


# ============================================================
# 3) Déduplication / classification des points
# ============================================================

def classify_points_type(points_df):
    df = points_df.copy()

    def _classify(row):
        fonc = row.get("TYPE_FONC")
        name = row.get("name") or ""
        if name.upper().startswith("SRO"):
            return "SRO"
        if fonc == "CHAMBRE":
            return "CHAMBRE"
        if fonc == "APPUI":
            return "POTEAU"
        if name == "JOINT":
            return "BPE"
        if name == "JOINT-SPLIT 1/8":
            return "BPE-PBO"
        if name == "SPLIT 1/8":
            return "PBO"
        return row.get("TYPE_STRUC") or fonc

    df["TYPE_OBJET"] = df.apply(_classify, axis=1)
    return df


def dedupe_exact_chambre_poteau(points_df, type_col="TYPE_OBJET", tolerance_m=0.5):
    """Fusionne les doublons CHAMBRE/POTEAU/SRO du même type situés à moins de
    tolerance_m mètres l'un de l'autre (peu importe le nom)."""
    df = points_df.copy()
    df["_lon"], df["_lat"] = zip(*df["coordinates"].map(_parse_lonlat))
    keep_mask = pd.Series(True, index=df.index)

    for t in ("CHAMBRE", "POTEAU", "SRO"):
        sub = df[df[type_col] == t]
        idxs = list(sub.index)
        visited = set()
        for i in range(len(idxs)):
            if idxs[i] in visited:
                continue
            for j in range(i + 1, len(idxs)):
                if idxs[j] in visited:
                    continue
                d = haversine_m(
                    sub.loc[idxs[i], "_lon"], sub.loc[idxs[i], "_lat"],
                    sub.loc[idxs[j], "_lon"], sub.loc[idxs[j], "_lat"],
                )
                if d <= tolerance_m:
                    keep_mask[idxs[j]] = False
                    visited.add(idxs[j])

    result = df[keep_mask].drop(columns=["_lon", "_lat"])
    return result.reset_index(drop=True)


# ============================================================
# 4) Rattachement câbles <-> BPE/PBO (Cable Entré / Cables Sortie)
# ============================================================

def dedupe_bpe_pbo_tolerance(bpe_pbo_df, name_col="NOM", tolerance_m=0.5):
    """Fusionne les doublons BPE/PBO/BPE-PBO/SRO : même NOM, coordonnées à
    moins de tolerance_m mètres l'une de l'autre (absorbe le bruit de
    précision flottante du KMZ, ex. 2 points identiques à 1cm près)."""
    df = bpe_pbo_df.copy()
    df["_lon"], df["_lat"] = zip(*df["coordinates"].map(_parse_lonlat))
    keep_mask = pd.Series(True, index=df.index)

    for nom, sub in df.groupby(name_col):
        idxs = list(sub.index)
        visited = set()
        for i in range(len(idxs)):
            if idxs[i] in visited:
                continue
            for j in range(i + 1, len(idxs)):
                if idxs[j] in visited:
                    continue
                d = haversine_m(
                    sub.loc[idxs[i], "_lon"], sub.loc[idxs[i], "_lat"],
                    sub.loc[idxs[j], "_lon"], sub.loc[idxs[j], "_lat"],
                )
                if d <= tolerance_m:
                    keep_mask[idxs[j]] = False
                    visited.add(idxs[j])

    result = df[keep_mask].drop(columns=["_lon", "_lat"])
    return result.reset_index(drop=True)


def add_cable_entree_sortie(points_df, lines_df, tolerance_m=1.0, cable_name_col="NOM",
                             type_col="TYPE_OBJET"):
    df = points_df.copy()
    df["_lon"], df["_lat"] = zip(*df["coordinates"].map(_parse_lonlat))

    cables = lines_df[lines_df[cable_name_col].notna() & (lines_df[cable_name_col] != "")]

    ref_lat = float(df["_lat"].mean())

    origin_lon, origin_lat, origin_cable = [], [], []
    end_lon, end_lat, end_cable = [], [], []
    for _, row in cables.iterrows():
        pts = row["coordinates"].strip().split()
        flon, flat = (float(v) for v in pts[0].split(",")[:2])
        llon, llat = (float(v) for v in pts[-1].split(",")[:2])
        origin_lon.append(flon); origin_lat.append(flat); origin_cable.append(row[cable_name_col])
        end_lon.append(llon); end_lat.append(llat); end_cable.append(row[cable_name_col])

    origin_idx = SpatialIndex(origin_lon, origin_lat, ref_lat=ref_lat)
    end_idx = SpatialIndex(end_lon, end_lat, ref_lat=ref_lat)

    entres, sorties = [], []
    for lon, lat, t in zip(df["_lon"], df["_lat"], df[type_col]):
        if t not in ("BPE", "PBO", "BPE-PBO", "SRO"):
            entres.append(""); sorties.append("")
            continue

        sortie_matches = [origin_cable[m] for m in origin_idx.query_ball_point(lon, lat, tolerance_m)]
        sorties.append("; ".join(sorted(set(sortie_matches))))

        if t == "SRO":
            entres.append("")
            continue

        entre_matches = [end_cable[m] for m in end_idx.query_ball_point(lon, lat, tolerance_m)]
        entres.append("; ".join(sorted(set(entre_matches))))

    df["Cable Entré"] = entres
    df["Cables Sortie"] = sorties
    df["Nbre_Cables_sorties"] = [len([x for x in s.split("; ") if x]) if s else 0 for s in sorties]
    return df.drop(columns=["_lon", "_lat"])


def _build_support_lookup(support_df, tolerance_m=1.0, type_col="TYPE_OBJET"):
    sup = support_df.copy()
    sup["_lon"], sup["_lat"] = zip(*sup["coordinates"].map(_parse_lonlat))
    ref_lat = float(sup["_lat"].mean()) if len(sup) else 0.0
    idx = SpatialIndex(sup["_lon"].tolist(), sup["_lat"].tolist(), ref_lat=ref_lat)
    sup_names = sup["name"].to_numpy()
    sup_types = sup[type_col].to_numpy()

    def lookup(lon, lat):
        return [(sup_names[m], sup_types[m]) for m in idx.query_ball_point(lon, lat, tolerance_m)]

    return lookup


def add_support_column(bpe_pbo_df, support_df, tolerance_m=1.0, type_col="TYPE_OBJET"):
    bpe = bpe_pbo_df.copy()
    bpe["_lon"], bpe["_lat"] = zip(*bpe["coordinates"].map(_parse_lonlat))
    lookup = _build_support_lookup(support_df, tolerance_m=tolerance_m, type_col=type_col)

    supports, types = [], []
    for lon, lat in zip(bpe["_lon"], bpe["_lat"]):
        matches = lookup(lon, lat)
        names_found, types_found = [], []
        for name, t in matches:
            names_found.append(name)
            types_found.append("Aerien" if t == "POTEAU" else ("SRO" if t == "SRO" else "GC"))
        supports.append("; ".join(names_found))
        types.append("; ".join(types_found))

    bpe = bpe.drop(columns=["_lon", "_lat"])
    bpe["Support"] = supports
    bpe["Type"] = types
    return bpe


# ============================================================
# 5) Fusion des câbles (segments -> 1 ligne par NOM) + Origine/Traversés
# ============================================================

def merge_cables(lines_df, support_df, tolerance_m=1.0, cable_name_col="NOM"):
    lookup = _build_support_lookup(support_df, tolerance_m=tolerance_m)
    cables = lines_df[lines_df[cable_name_col].notna() & (lines_df[cable_name_col] != "")]

    rows = []
    for nom, group in cables.groupby(cable_name_col):
        segments = []
        for _, r in group.iterrows():
            pts = r["coordinates"].strip().split()
            segments.append([tuple(float(v) for v in p.split(",")[:2]) for p in pts])

        last_coords = [seg[-1] for seg in segments]
        all_vertices = [v for seg in segments for v in seg]

        origin_candidates = []
        for seg in segments:
            first = seg[0]
            if not any(_coords_close(first, lc, tolerance_m) for lc in last_coords if lc is not first):
                origin_candidates.append(first)
        true_origin = origin_candidates[0] if origin_candidates else segments[0][0]

        origine_name = "; ".join(sorted(set(n for n, t in lookup(*true_origin))))

        traversed_chambre, traversed_poteau = set(), set()
        for v in all_vertices:
            if _coords_close(v, true_origin, tolerance_m):
                continue
            for name, t in lookup(v[0], v[1]):
                if t == "CHAMBRE":
                    traversed_chambre.add(name)
                elif t == "POTEAU":
                    traversed_poteau.add(name)

        longueur_totale = 0.0
        for seg in segments:
            for a, b in zip(seg[:-1], seg[1:]):
                longueur_totale += haversine_m(a[0], a[1], b[0], b[1])

        nb_ch, nb_pot = len(traversed_chambre), len(traversed_poteau)

        base = group.iloc[0].to_dict()
        base[cable_name_col] = nom
        base["ORIGINE"] = origine_name
        base["Supports_Traversés"] = "; ".join(sorted(traversed_chambre | traversed_poteau))
        base["Nbre_CH_Traversée"] = nb_ch
        base["Nbre_POT_Traversée"] = nb_pot
        base["LONGUEUR"] = round(longueur_totale, 2)
        base["LOG_Total"] = round(longueur_totale + nb_ch * 30 + nb_pot * 10, 2)
        rows.append(base)

    result = pd.DataFrame(rows).reset_index(drop=True)
    return result.drop(columns=["name", "geometry_type", "coordinates", "EXTREMITE"], errors="ignore")


# ============================================================
# 6) Pipeline complet (BPE_PBO, Support, Cable (Ligne))
# ============================================================

def build_all_sheets(kmz_path, tolerance_m=1.0, points_df=None, lines_df=None, pre_classified_df=None):
    """
    Si pre_classified_df est fourni (points_df déjà classé, colonne TYPE_OBJET
    déjà présente, ex. via l'outil "Générer" basé sur generic_classifier),
    on saute le parsing/la classification codée en dur et on l'utilise tel
    quel. Sinon comportement inchangé (classify_points_type codé en dur).
    """
    if points_df is None or lines_df is None:
        points_df, lines_df = build_layers(kmz_path)
    if pre_classified_df is not None:
        df = pre_classified_df
    else:
        df = classify_points_type(points_df)
    df = dedupe_exact_chambre_poteau(df, type_col="TYPE_OBJET")

    support_df = df[df["TYPE_OBJET"].isin(["CHAMBRE", "POTEAU"])].reset_index(drop=True)
    bpe_pbo_df = df[df["TYPE_OBJET"].isin(["BPE", "PBO", "BPE-PBO", "SRO"])].reset_index(drop=True)

    bpe_pbo_df = add_cable_entree_sortie(bpe_pbo_df, lines_df, tolerance_m=tolerance_m)
    bpe_pbo_df = dedupe_bpe_pbo_tolerance(bpe_pbo_df, tolerance_m=0.5)
    bpe_pbo_df = add_support_column(bpe_pbo_df, support_df, tolerance_m=tolerance_m)

    support_for_cable_origin = df[df["TYPE_OBJET"].isin(["CHAMBRE", "POTEAU", "SRO"])].reset_index(drop=True)
    cables_df = merge_cables(lines_df, support_for_cable_origin, tolerance_m=tolerance_m)

    bpe_pbo_export = bpe_pbo_df.drop(columns=["name", "geometry_type"])
    support_export = support_df.drop(columns=["geometry_type"]) if "geometry_type" in support_df.columns else support_df

    return {
        "points_df": points_df, "lines_df": lines_df, "df": df,
        "support_df": support_df, "bpe_pbo_df": bpe_pbo_df, "cables_df": cables_df,
        "bpe_pbo_export": bpe_pbo_export, "support_export": support_export,
    }


# ============================================================
# 7) Synoptique : hiérarchie des câbles + rendu Excel
# ============================================================

def parse_cable_name(nom):
    m = re.match(r"^CDI_(.+)[_-]([^_-]+)$", nom)
    if not m:
        return None, None
    path_str, capacite = m.group(1), m.group(2)
    return tuple(re.split(r"[_-]", path_str)), capacite


def _cable_true_origin(group, tolerance_m):
    """Calcule l'origine géométrique réelle d'un câble (le sommet qui n'est
    le point d'arrivée d'aucun autre segment du même câble)."""
    segments = []
    for _, r in group.iterrows():
        pts = r["coordinates"].strip().split()
        segments.append([tuple(float(v) for v in p.split(",")[:2]) for p in pts])

    last_coords = [seg[-1] for seg in segments]
    origin_candidates = []
    for seg in segments:
        first = seg[0]
        if not any(_coords_close(first, lc, tolerance_m) for lc in last_coords if lc is not first):
            origin_candidates.append(first)
    return origin_candidates[0] if origin_candidates else segments[0][0]


def find_nearest_point(lon, lat, ref_df):
    """Recherche brute-force du point le plus proche dans ref_df (coordinates)."""
    best_idx, best_dist = None, float("inf")
    for idx, row in ref_df.iterrows():
        rlon, rlat = _parse_lonlat(row["coordinates"])
        if rlon is None:
            continue
        d = haversine_m(lon, lat, rlon, rlat)
        if d < best_dist:
            best_dist, best_idx = d, idx
    return best_idx, best_dist


def compute_joint_passing_cables(lines_df, bpe_pbo_df, tolerance_m=1.0, cable_name_col="NOM"):
    """Pour chaque point BPE/PBO/SRO, l'ensemble des câbles dont le tracé
    passe à proximité (tous les sommets, pas seulement les extrémités) —
    sert de repli quand 'Cable Entré' est vide (câble parent qui continue
    au-delà de ce point plutôt que de s'y terminer littéralement)."""
    bpe = bpe_pbo_df.copy()
    bpe["_lon"], bpe["_lat"] = zip(*bpe["coordinates"].map(_parse_lonlat))
    ref_lat = float(bpe["_lat"].mean()) if len(bpe) else 0.0
    m_per_deg_lat = 111320.0
    m_per_deg_lon = 111320.0 * math.cos(math.radians(ref_lat))
    bx = (bpe["_lon"] * m_per_deg_lon).to_numpy()
    by = (bpe["_lat"] * m_per_deg_lat).to_numpy()
    b_idx = bpe.index.to_numpy()

    result = {}
    cables = lines_df[lines_df[cable_name_col].notna() & (lines_df[cable_name_col] != "")]
    for nom, group in cables.groupby(cable_name_col):
        chain = _order_segments(
            [[tuple(float(v) for v in p.split(",")[:2]) for p in r["coordinates"].strip().split()]
             for _, r in group.iterrows()],
            tolerance_m)
        for lon, lat in chain:
            x, y = lon * m_per_deg_lon, lat * m_per_deg_lat
            d = np.hypot(bx - x, by - y)
            for m in np.where(d <= tolerance_m)[0]:
                idx = b_idx[m]
                result.setdefault(idx, set()).add(nom)
    return result


def build_cable_hierarchy_geometric(lines_df, bpe_pbo_df, tolerance_m=1.0, warn_threshold_m=2.0,
                                     cable_name_col="NOM", joint_passing_cables=None):
    """
    Construit la hiérarchie des câbles par GÉOMÉTRIE (pas par nom) :
    l'origine réelle de chaque câble est rapprochée du SRO/BPE/BPE-PBO le plus
    proche (SRO/BPE/BPE-PBO/PBO : un PBO reste un point de rattachement valide,
    l'incohérence "PBO utilisé comme BPE" est détectée séparément par la
    règle 3, find_pbo_type_anomalies, sans fausser la distance de rattachement). Si c'est le SRO -> câble racine. Sinon -> le câble
    parent est celui du 'Cable Entré' de ce joint ; si ce champ est vide (le
    parent continue au-delà de ce point plutôt que de s'y terminer), on
    cherche quel(s) autre(s) câble(s) passent réellement par ce point (repli).
    2 joints différents dans une même chambre restent 2 points distincts,
    donc 2 familles de câbles séparées.
    Si la distance au point le plus proche dépasse warn_threshold_m, le
    rattachement est quand même fait (au plus proche) mais marqué suspect.
    """
    cables = lines_df[lines_df[cable_name_col].notna() & (lines_df[cable_name_col] != "")]
    if joint_passing_cables is None:
        joint_passing_cables = compute_joint_passing_cables(lines_df, bpe_pbo_df, tolerance_m)

    children_of = {}
    roots = []
    attach_of = {}      # nom câble -> index bpe_pbo_df du point de rattachement
    suspects = {}        # nom câble -> (distance_m, nom_du_joint_le_plus_proche, type_anomalie)

    for nom, group in cables.groupby(cable_name_col):
        true_origin = _cable_true_origin(group, tolerance_m)
        nearest_idx, dist = find_nearest_point(true_origin[0], true_origin[1], bpe_pbo_df)

        # Un PBO proche (<=warn_threshold_m) est un rattachement valide (ex. un
        # PBO directement raccordé au câble racine). Mais si ce point le plus
        # proche est déjà "loin" (> warn_threshold_m, câble suspect), on
        # cherche spécifiquement le SRO/BPE/BPE-PBO le plus proche à la place
        # d'un PBO seul (un PBO n'est jamais un vrai "tronc" plausible dans ce
        # cas d'incertitude).
        if dist > warn_threshold_m:
            bpe_only_df = bpe_pbo_df[bpe_pbo_df["TYPE_OBJET"].isin(["SRO", "BPE", "BPE-PBO"])]
            bpe_idx, bpe_dist = find_nearest_point(true_origin[0], true_origin[1], bpe_only_df)
            if bpe_idx is not None:
                nearest_idx, dist = bpe_idx, bpe_dist

        attach_of[nom] = nearest_idx

        if nearest_idx is None:
            roots.append(nom)
            continue

        if dist > warn_threshold_m:
            joint_label = bpe_pbo_df.loc[nearest_idx, "Support"] or bpe_pbo_df.loc[nearest_idx, "NOM"]
            suspects[nom] = (round(dist, 2), joint_label, "Distance excessive")

        if bpe_pbo_df.loc[nearest_idx, "TYPE_OBJET"] == "SRO":
            roots.append(nom)
            continue

        parent_nom = bpe_pbo_df.loc[nearest_idx, "Cable Entré"]
        if not parent_nom:
            candidates = joint_passing_cables.get(nearest_idx, set()) - {nom}
            parent_nom = sorted(candidates)[0] if candidates else ""

        if parent_nom:
            children_of.setdefault(parent_nom, []).append(nom)
        else:
            roots.append(nom)  # aucun parent identifiable -> racine par défaut (à revoir)
            if nom not in suspects:
                joint_label = bpe_pbo_df.loc[nearest_idx, "Support"] or bpe_pbo_df.loc[nearest_idx, "NOM"]
                suspects[nom] = (round(dist, 2), joint_label, "Aucun câble parent trouvé")

    return children_of, roots, attach_of, suspects


def find_pbo_type_anomalies(bpe_pbo_df):
    """
    Règle 3 : une boîte classée PBO seul (ni BPE, ni BPE-PBO) ne devrait
    normalement avoir qu'UN SEUL câble de distribution en sortie cataloguée
    (elle alimente des abonnés, pas d'autres câbles de distribution). Si elle
    en a 2 ou plus, c'est une incohérence probable de conception/étude (la
    boîte devrait être un BPE ou BPE-PBO).
    Retourne une liste de (index bpe_pbo_df, nom du joint, nb de câbles sortie).
    """
    anomalies = []
    mask = (bpe_pbo_df["TYPE_OBJET"] == "PBO") & (bpe_pbo_df["Nbre_Cables_sorties"] >= 2)
    for idx, row in bpe_pbo_df[mask].iterrows():
        joint_label = row.get("Support") or row.get("NOM")
        anomalies.append((idx, joint_label, int(row["Nbre_Cables_sorties"])))
    return anomalies


def resolve_attach_positions(children_of, chains, attach_of):
    """Pour chaque câble parent, regroupe ses enfants par POSITION dans sa
    propre liste de joints (chains[parent]['joints']), à partir de
    attach_of (index bpe_pbo_df déjà déterminé par la géométrie)."""
    resolved = {}
    for parent, kids in children_of.items():
        joints = chains.get(parent, {}).get("joints", [])
        groups = {}
        for k in kids:
            j_idx = attach_of.get(k)
            if joints and j_idx in joints:
                pos = joints.index(j_idx)
            else:
                pos = len(joints) - 1 if joints else 0
            groups.setdefault(pos, []).append(k)
        resolved[parent] = groups
    return resolved


def reorder_children_by_position(children_of, resolved_groups):
    """Trie les enfants par position d'attache DÉCROISSANTE (dernier joint du
    parent traité en premier, cf. méthode validée)."""
    for nom, kids in children_of.items():
        groups = resolved_groups.get(nom, {})
        pos_of = {}
        for pos, klist in groups.items():
            for k in klist:
                pos_of[k] = pos
        kids.sort(key=lambda k: -pos_of.get(k, 0))


def build_cable_hierarchy(cables_df):
    path_to_nom = {}
    for _, row in cables_df.iterrows():
        segs, cap = parse_cable_name(row["NOM"])
        if segs:
            path_to_nom[segs] = row["NOM"]

    children_of = {}
    roots = []
    for segs, nom in path_to_nom.items():
        if len(segs) == 1:
            roots.append(nom)
        else:
            parent_nom = path_to_nom.get(segs[:-1])
            if parent_nom:
                children_of.setdefault(parent_nom, []).append(nom)
            else:
                roots.append(nom)
    return children_of, roots


def format_box_label(row):
    t = row.get("TYPE_OBJET", "")
    nom = row.get("NOM", "")
    coords = row.get("coordinates", "")
    sortie = row.get("Cables Sortie", "") or ""
    sortie_names = [x.strip() for x in sortie.split(";") if x.strip()]
    sortie_lines = "\n".join(f"Cable sortie : {c}" for c in sortie_names)

    if t == "SRO":
        base = f"TYPE_OBJET : {t}\nNOM : {nom}\nCoordinates : {coords}"
    else:
        base = (f"TYPE_OBJET : {t}\nNOM : {nom}\nSupport : {row.get('Support','')}\n"
                f"Type : {row.get('Type','')}\nCoordinates : {coords}")

    return f"{base}\n{sortie_lines}" if sortie_lines else base


def _order_segments(segments, tol):
    last_coords = [seg[-1] for seg in segments]
    start_idx = 0
    for i, seg in enumerate(segments):
        if not any(_coords_close(seg[0], lc, tol) for j, lc in enumerate(last_coords) if j != i):
            start_idx = i
            break
    used = {start_idx}
    chain = list(segments[start_idx])
    changed = True
    while changed:
        changed = False
        current_end = chain[-1]
        for i, seg in enumerate(segments):
            if i in used:
                continue
            if _coords_close(seg[0], current_end, tol):
                chain.extend(seg[1:]); used.add(i); changed = True; break
            if _coords_close(seg[-1], current_end, tol):
                chain.extend(list(reversed(seg))[1:]); used.add(i); changed = True; break
    for i, seg in enumerate(segments):
        if i not in used:
            chain.extend(seg)
    return chain


def compute_cable_chains(lines_df, bpe_pbo_df, support_df, tolerance_m=1.0, cable_name_col="NOM",
                          extra_joints_by_cable=None):
    bpe = bpe_pbo_df[bpe_pbo_df["TYPE_OBJET"] != "SRO"].copy()
    bpe["_lon"], bpe["_lat"] = zip(*bpe["coordinates"].map(_parse_lonlat))
    ref_lat = float(bpe["_lat"].mean()) if len(bpe) else 0.0
    b_index = SpatialIndex(bpe["_lon"].tolist(), bpe["_lat"].tolist(), ref_lat=ref_lat)
    b_idx = bpe.index.to_numpy()
    b_entree = bpe["Cable Entré"].fillna("").to_numpy()
    b_sortie = bpe["Cables Sortie"].fillna("").to_numpy()

    extra_joints_by_cable = extra_joints_by_cable or {}

    support_lookup = _build_support_lookup(support_df, tolerance_m=tolerance_m)

    cables = lines_df[lines_df[cable_name_col].notna() & (lines_df[cable_name_col] != "")]
    result = {}
    for nom, group in cables.groupby(cable_name_col):
        segments = []
        for _, r in group.iterrows():
            pts = r["coordinates"].strip().split()
            segments.append([tuple(float(v) for v in p.split(",")[:2]) for p in pts])
        chain = _order_segments(segments, tolerance_m)
        my_extra_joints = extra_joints_by_cable.get(nom, set())

        joints, seg_supports, seg_lengths = [], [], []
        current = set()
        current_len = 0.0
        prev_vertex = None
        for lon, lat in chain:
            if prev_vertex is not None:
                current_len += haversine_m(prev_vertex[0], prev_vertex[1], lon, lat)
            prev_vertex = (lon, lat)

            for name, t in support_lookup(lon, lat):
                if t in ("CHAMBRE", "POTEAU"):
                    current.add(name)
            for m in b_index.query_ball_point(lon, lat, tolerance_m):
                # Un même point géométrique peut être partagé par plusieurs câbles
                # (conduite/chambre commune) sans que CE câble n'y soit réellement
                # raccordé -> on ne garde le joint que s'il a un lien logique réel
                # avec 'nom' (son extrémité, ou l'un de ses câbles sortants), OU
                # s'il s'agit d'un point de rattachement déjà confirmé par la
                # hiérarchie géométrique (extra_joints_by_cable) — seule source de
                # vérité en cas d'ambiguïté (plusieurs câbles passant exactement
                # par le même point).
                idx = b_idx[m]
                sortie_names = [x.strip() for x in b_sortie[m].split(";") if x.strip()]
                if b_entree[m] != nom and nom not in sortie_names and idx not in my_extra_joints:
                    continue
                if not joints or joints[-1] != idx:
                    joints.append(idx)
                    seg_supports.append(current)
                    seg_lengths.append(round(current_len, 2))
                    current = set()
                    current_len = 0.0
        result[nom] = {"joints": joints, "segment_supports": seg_supports, "segment_lengths": seg_lengths}
    return result


def find_attach_index(parent_nom, child_nom, chains, bpe_pbo_df):
    joints = chains.get(parent_nom, {}).get("joints", [])
    for i, jidx in enumerate(joints):
        sortie_val = bpe_pbo_df.loc[jidx, "Cables Sortie"] or ""
        names = [x.strip() for x in sortie_val.split(";") if x.strip()]
        if child_nom in names:
            return i
    return len(joints) - 1 if joints else None


def reorder_children_by_joint(children_of, chains, bpe_pbo_df):
    for nom, kids in children_of.items():
        attach = {k: find_attach_index(nom, k, chains, bpe_pbo_df) for k in kids}
        kids.sort(key=lambda k: -(attach[k] if attach[k] is not None else -1))


def layout_rows(children_of, roots):
    row_pos = {}
    next_row = [0]

    def assign(nom):
        row_pos[nom] = next_row[0]
        next_row[0] += 1
        for k in children_of.get(nom, []):
            assign(k)

    for r in roots:
        assign(r)
    return row_pos


def render_synoptique(bpe_pbo_df, sro_idx, roots, children_of, row_pos, chains, cable_info,
                       resolved_groups, suspects, attach_of, out_path=None, wb=None):
    if wb is None:
        wb = Workbook()
        ws = wb.active
    else:
        ws = wb.create_sheet("Synoptique")
    ws.title = "Synoptique"

    suspect_attach_joints = {attach_of[c] for c in suspects if attach_of.get(c) is not None}
    suspects_by_joint = {}
    for c in suspects:
        j = attach_of.get(c)
        if j is not None:
            suspects_by_joint.setdefault(j, []).append(c)
    type_incoherent_joints = {
        idx for idx, _, _ in find_pbo_type_anomalies(bpe_pbo_df)
    }

    BOX_COLS = 4
    BOX_ROWS = 13
    CABLE_COLS = 3
    ROW_SCALE = 15

    thin = Side(style="thin", color="000000")
    box_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    red_side = Side(style="thin", color="FF0000")
    red_border = Border(left=red_side, right=red_side, top=red_side, bottom=red_side)
    red_font = Font(color="FF0000", bold=True, size=9)
    header_font = Font(color="C55A11", bold=True, size=10)
    label_font = Font(color="1F4E78", bold=True, size=9)
    value_font = Font(color="000000", size=9)
    green_font = Font(color="00863B", bold=True, size=9)

    FILL_GC = "FDE9D9"       # orangé clair
    FILL_AERIEN = "D9EAD3"   # vert clair
    LINE_CONDUITE = "C55A11"  # orangé foncé
    LINE_AERIEN = "375623"    # vert foncé

    def _box_colors(type_val):
        t = (type_val or "").upper()
        if "AERIEN" in t:
            return PatternFill(fill_type="solid", start_color=FILL_AERIEN, end_color=FILL_AERIEN)
        return PatternFill(fill_type="solid", start_color=FILL_GC, end_color=FILL_GC)

    def _cable_line_color(type_struc):
        t = (type_struc or "").upper()
        if "AERIEN" in t:
            return LINE_AERIEN
        return LINE_CONDUITE

    def row_start(r):
        return 2 + round(r * ROW_SCALE)

    def _merge_cell(r0, c0, r1, c1, value, font, align="left", border=None):
        ws.merge_cells(start_row=r0, start_column=c0, end_row=r1, end_column=c1)
        cell = ws.cell(row=r0, column=c0)
        cell.value = value
        cell.font = font
        cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
        b = border if border is not None else box_border
        for rr in range(r0, r1 + 1):
            for cc in range(c0, c1 + 1):
                ws.cell(row=rr, column=cc).border = b

    def draw_box(r0, c0, row, cum_length, joint_idx=None):
        """Boîte de taille fixe 4 colonnes x 13 lignes, remplissage uniforme
        selon le Type (GC=orangé clair, Aerien=vert clair) : Support / longueur
        SRO->Joint, NOM, TYPE_OBJET / Type, Coordinates, tableau Cable entré /
        Cables Sorties (1 ligne par câble sortant), padding jusqu'à 13 lignes.
        Bordure rouge si cette boîte est le point de rattachement d'un câble
        suspect (règles 1/2), ou si elle est elle-même une incohérence de
        type (règle 3, PBO seul avec >=2 câbles sortie) ; les câbles suspects
        dans la liste 'Cables Sorties' sont affichés en rouge (règles 1/2)."""
        half = BOX_COLS // 2
        c_mid = c0 + half
        c_end = c0 + BOX_COLS - 1
        r_end = r0 + BOX_ROWS - 1

        is_flagged = joint_idx is not None and (
            joint_idx in suspect_attach_joints or joint_idx in type_incoherent_joints)
        this_border = red_border if is_flagged else box_border

        fill = _box_colors(row.get("Type", ""))
        for rr in range(r0, r_end + 1):
            for cc in range(c0, c_end + 1):
                cell = ws.cell(row=rr, column=cc)
                cell.fill = fill
                cell.border = this_border

        support_label = row.get("Support") or row.get("TYPE_OBJET") or row.get("NOM") or ""
        entree = (row.get("Cable Entré") or "").strip()
        sortie_list = [x.strip() for x in (row.get("Cables Sortie") or "").split(";") if x.strip()]
        for c in suspects_by_joint.get(joint_idx, []):
            if c not in sortie_list:
                sortie_list.append(c)
        nb_lignes_cable = max(1, len(sortie_list))

        r = r0
        _merge_cell(r, c0, r, c_mid - 1, support_label, header_font, border=this_border)
        _merge_cell(r, c_mid, r, c_end, f"{cum_length:.2f} m" if cum_length is not None else "",
                    header_font, align="center", border=this_border)
        r += 1
        _merge_cell(r, c0, r, c_end, row.get("NOM", ""), green_font, align="center", border=this_border)
        r += 1
        _merge_cell(r, c0, r, c_mid - 1, row.get("TYPE_OBJET", ""), label_font, border=this_border)
        _merge_cell(r, c_mid, r, c_end, row.get("Type", ""), value_font, border=this_border)
        r += 1
        _merge_cell(r, c0, r, c_end, row.get("coordinates", ""), value_font, border=this_border)
        r += 1
        _merge_cell(r, c0, r, c_mid - 1, "Cable entré", label_font, border=this_border)
        _merge_cell(r, c_mid, r, c_end, "Cables Sorties", label_font, border=this_border)
        r += 1
        for i in range(nb_lignes_cable):
            if r > r_end:
                break
            sortie_val = sortie_list[i] if i < len(sortie_list) else ""
            sortie_font = red_font if sortie_val in suspects else value_font
            _merge_cell(r, c0, r, c_mid - 1, entree, value_font, border=this_border)
            _merge_cell(r, c_mid, r, c_end, sortie_val, sortie_font, border=this_border)
            r += 1

    def draw_segment(row_line, col_from, col_to, cable_nom, type_struc, supports, seg_length):
        is_suspect = cable_nom in suspects
        line_color = "FF0000" if is_suspect else _cable_line_color(type_struc)
        seg_side = Side(style="medium", color=line_color)
        seg_font = Font(color=line_color, size=9, bold=True)
        for cc in range(col_from, col_to + 1):
            existing_left = ws.cell(row=row_line, column=cc).border.left
            left = existing_left if (existing_left and existing_left.style) else None
            if cc == col_from:
                left = seg_side
            ws.cell(row=row_line, column=cc).border = Border(left=left, bottom=seg_side)

        suffix = "  /!\\ RATTACHEMENT SUSPECT" if is_suspect else ""
        value_cell = ws.cell(row=row_line - 1, column=col_from)
        value_cell.value = f"{cable_nom} - {type_struc} - {seg_length} m{suffix}"
        value_cell.font = seg_font

        c3 = ws.cell(row=row_line + 1, column=col_from)
        c3.value = f"Supports_Traversés : {'; '.join(sorted(supports))}"
        c3.font = value_font

    def fan_out(trunk_col, origin_row, row_lines):
        all_rows = row_lines + [origin_row]
        rmin, rmax = min(all_rows), max(all_rows)
        if rmin == rmax:
            return
        for rr in range(rmin, rmax + 1):
            existing = ws.cell(row=rr, column=trunk_col).border
            ws.cell(row=rr, column=trunk_col).border = Border(
                left=thin, bottom=existing.bottom if existing else None)

    sro_row_line = (row_start(row_pos.get(roots[0], 0)) + BOX_ROWS // 2) if roots else (row_start(0) + BOX_ROWS // 2)
    sro_col = 1
    sro_row_data = bpe_pbo_df.loc[sro_idx] if sro_idx is not None else pd.Series({"TYPE_OBJET": "SRO"})
    draw_box(row_start(row_pos.get(roots[0], 0)) if roots else row_start(0), sro_col, sro_row_data, 0.0, joint_idx=sro_idx)

    def render(nom, start_col, row_line, is_root=False, cum_length_start=0.0):
        info = chains.get(nom, {"joints": [], "segment_supports": [], "segment_lengths": []})
        joints, seg_supports = info["joints"], info["segment_supports"]
        seg_lengths = info.get("segment_lengths", [0.0] * len(joints))

        skip_first = (not is_root) and len(joints) > 1

        col = start_col
        joint_box_col = [start_col] if skip_first else []
        cinfo = cable_info.get(nom)
        type_struc = cinfo.get("TYPE_STRUC") if cinfo is not None else ""

        cum = cum_length_start
        joint_cum = []
        for i, jidx in enumerate(joints):
            cum += seg_lengths[i] if i < len(seg_lengths) else 0.0
            if skip_first and i == 0:
                joint_cum.append(cum)
                continue
            box_col = col + CABLE_COLS
            draw_segment(row_line, col, box_col - 1, nom, type_struc,
                         seg_supports[i], seg_lengths[i] if i < len(seg_lengths) else 0.0)
            draw_box(row_start(row_pos[nom]), box_col, bpe_pbo_df.loc[jidx], cum, joint_idx=jidx)
            joint_box_col.append(box_col)
            joint_cum.append(cum)
            col = box_col + BOX_COLS

        kids = children_of.get(nom, [])
        if not kids or not joints:
            return

        groups = resolved_groups.get(nom, {})
        order = []
        for k in kids:
            for pos, klist in groups.items():
                if k in klist and pos not in order:
                    order.append(pos)
                    break

        for a in order:
            group_kids = groups[a]
            trunk_col = joint_box_col[a] + BOX_COLS if a < len(joint_box_col) else joint_box_col[-1] + BOX_COLS
            rows_lines = [row_start(row_pos[k]) + BOX_ROWS // 2 for k in group_kids]
            fan_out(trunk_col, row_line, rows_lines)
            for k, rl in zip(group_kids, rows_lines):
                render(k, trunk_col, rl, cum_length_start=joint_cum[a] if a < len(joint_cum) else cum)

    if roots:
        trunk_col = sro_col + BOX_COLS
        rows_lines = [row_start(row_pos[r]) + BOX_ROWS // 2 for r in roots]
        fan_out(trunk_col, sro_row_line, rows_lines)
        for r, rl in zip(roots, rows_lines):
            render(r, trunk_col, rl, is_root=True, cum_length_start=0.0)

    for c in range(1, 300):
        ws.column_dimensions[get_column_letter(c)].width = 10

    if out_path:
        wb.save(out_path)
    return wb


# ============================================================
# 8) Export global (4 onglets) — point d'entrée principal
# ============================================================

def build_global_excel(kmz_path, out_path, tolerance_m=1.0, progress_cb=None,
                        points_df=None, lines_df=None, pre_classified_df=None):
    """
    Génère le fichier Excel complet (BPE_PBO, Support, Cable (Ligne), Synoptique)
    à partir d'un fichier KMZ. progress_cb(pct, message) est appelé pour le suivi
    (optionnel, pour affichage de progression dans QGIS). pre_classified_df
    permet d'utiliser une classification déjà faite (outil "Générer").
    """
    def _p(pct, msg):
        if progress_cb:
            progress_cb(pct, msg)

    _p(5, "Lecture du KMZ...")
    sheets = build_all_sheets(kmz_path, tolerance_m=tolerance_m, points_df=points_df,
                               lines_df=lines_df, pre_classified_df=pre_classified_df)

    _p(45, "Écriture des feuilles BPE_PBO / Support / Cable...")
    wb = Workbook()
    ws0 = wb.active
    wb.remove(ws0)

    from openpyxl.utils.dataframe import dataframe_to_rows

    def _write_df(wb, sheet_name, df):
        ws = wb.create_sheet(sheet_name)
        for r in dataframe_to_rows(df, index=False, header=True):
            ws.append(r)

    _write_df(wb, "BPE_PBO", sheets["bpe_pbo_export"])
    _write_df(wb, "Support", sheets["support_export"])
    _write_df(wb, "Cable (Ligne)", sheets["cables_df"])

    _p(60, "Construction de la hiérarchie des câbles...")
    cables_df = sheets["cables_df"]
    bpe_pbo_df = sheets["bpe_pbo_df"]
    lines_df = sheets["lines_df"]
    support_df = sheets["support_df"]

    cable_info = {r["NOM"]: r for _, r in cables_df.iterrows()}

    _p(65, "Calcul du passage géométrique des câbles...")
    joint_passing_cables = compute_joint_passing_cables(lines_df, bpe_pbo_df, tolerance_m=tolerance_m)

    _p(68, "Construction de la hiérarchie des câbles (géométrie)...")
    children_of, roots, attach_of, suspects = build_cable_hierarchy_geometric(
        lines_df, bpe_pbo_df, tolerance_m=tolerance_m, warn_threshold_m=2.0,
        joint_passing_cables=joint_passing_cables)

    extra_joints_by_cable = {}
    for parent, kids in children_of.items():
        pts = {attach_of[k] for k in kids if attach_of.get(k) is not None}
        if pts:
            extra_joints_by_cable[parent] = pts

    _p(72, "Calcul des joints le long de chaque câble...")
    chains = compute_cable_chains(lines_df, bpe_pbo_df, support_df, tolerance_m=tolerance_m,
                                   extra_joints_by_cable=extra_joints_by_cable)

    resolved_groups = resolve_attach_positions(children_of, chains, attach_of)
    reorder_children_by_position(children_of, resolved_groups)
    row_pos = layout_rows(children_of, roots)

    sro_idx_list = bpe_pbo_df.index[bpe_pbo_df["TYPE_OBJET"] == "SRO"].tolist()
    sro_idx = sro_idx_list[0] if sro_idx_list else None

    pbo_type_anomalies = find_pbo_type_anomalies(bpe_pbo_df)

    if suspects or pbo_type_anomalies:
        _p(78, f"{len(suspects)} câble(s) et {len(pbo_type_anomalies)} boîte(s) signalé(s)...")
        anomalies_ws = wb.create_sheet("Anomalies")
        anomalies_ws.append(["Type d'anomalie", "Câble / Boîte", "Distance (m)", "Détail"])
        for nom, (dist, joint_label, anomaly_type) in suspects.items():
            anomalies_ws.append([anomaly_type, nom, dist, f"Joint le plus proche : {joint_label}"])
        for idx, joint_label, nb_sorties in pbo_type_anomalies:
            anomalies_ws.append(["Incohérence de type (PBO)", joint_label, "",
                                  f"PBO avec {nb_sorties} câbles sortie (devrait être BPE/BPE-PBO)"])

    _p(85, "Génération du synoptique...")
    render_synoptique(bpe_pbo_df, sro_idx, roots, children_of, row_pos, chains, cable_info,
                       resolved_groups, suspects, attach_of, out_path=None, wb=wb)

    _p(95, "Sauvegarde du fichier Excel...")
    wb.save(out_path)
    _p(100, "Terminé.")
    return out_path


# ============================================================
# 9) Algorithme QGIS Processing
# ============================================================
try:
    from qgis.core import (
        QgsProcessingAlgorithm,
        QgsProcessingParameterFile,
        QgsProcessingParameterNumber,
        QgsProcessingParameterFileDestination,
        QgsProcessingException,
    )
    from qgis.PyQt.QtCore import QCoreApplication
    _QGIS_AVAILABLE = True
except ImportError:
    _QGIS_AVAILABLE = False


if _QGIS_AVAILABLE:

    class FTTHApdKmzToExcelAlgorithm(QgsProcessingAlgorithm):
        """
        Extrait un plan de câblage FTTH (KMZ) vers un classeur Excel complet
        (BPE_PBO, Support, Cable (Ligne), Synoptique).
        """

        INPUT_KMZ = "INPUT_KMZ"
        TOLERANCE = "TOLERANCE"
        OUTPUT_XLSX = "OUTPUT_XLSX"

        def tr(self, string):
            return QCoreApplication.translate("FTTHApdKmzToExcelAlgorithm", string)

        def createInstance(self):
            return FTTHApdKmzToExcelAlgorithm()

        def name(self):
            return "ftth_apd_kmz_to_excel"

        def displayName(self):
            return self.tr("APD FTTH - KMZ vers Excel (BPE_PBO, Support, Câbles, Synoptique)")

        def group(self):
            return self.tr("FTTH")

        def groupId(self):
            return "ftth"

        def shortHelpString(self):
            return self.tr(
                "Extrait un plan de câblage FTTH au format KMZ vers un classeur Excel "
                "complet à 4 onglets :\n"
                "- BPE_PBO : boîtiers (BPE/PBO/SRO) avec câbles entrant/sortant\n"
                "- Support : chambres et poteaux\n"
                "- Cable (Ligne) : câbles fusionnés avec origine, longueur, supports traversés\n"
                "- Synoptique : représentation graphique de l'arborescence du réseau\n\n"
                "La 'Tolérance' définit la distance (en mètres) en dessous de laquelle deux "
                "points sont considérés comme confondus (ex. un câble qui s'arrête à un "
                "boîtier). 1 mètre convient dans la grande majorité des cas."
            )

        def initAlgorithm(self, config=None):
            self.addParameter(
                QgsProcessingParameterFile(
                    self.INPUT_KMZ,
                    self.tr("Fichier KMZ du plan de câblage"),
                    extension="kmz",
                )
            )
            self.addParameter(
                QgsProcessingParameterNumber(
                    self.TOLERANCE,
                    self.tr("Tolérance de rapprochement des points (mètres)"),
                    type=QgsProcessingParameterNumber.Double,
                    defaultValue=1.0,
                    minValue=0.01,
                    maxValue=50.0,
                )
            )
            self.addParameter(
                QgsProcessingParameterFileDestination(
                    self.OUTPUT_XLSX,
                    self.tr(
                        "Fichier Excel de sortie (par défaut : même dossier que le KMZ, "
                        "APD_FTTH_GLOBAL.xlsx)"
                    ),
                    fileFilter="Fichiers Excel (*.xlsx)",
                    optional=True,
                )
            )

        def processAlgorithm(self, parameters, context, feedback):
            kmz_path = self.parameterAsFile(parameters, self.INPUT_KMZ, context)
            tolerance = self.parameterAsDouble(parameters, self.TOLERANCE, context)
            out_path = self.parameterAsFileOutput(parameters, self.OUTPUT_XLSX, context)

            if not out_path:
                # Par défaut : même dossier que le KMZ importé
                out_path = os.path.join(os.path.dirname(kmz_path), "APD_FTTH_GLOBAL.xlsx")

            if not out_path.lower().endswith(".xlsx"):
                out_path += ".xlsx"

            missing = []
            for mod in ("lxml", "pandas", "numpy", "openpyxl"):
                try:
                    __import__(mod)
                except ImportError:
                    missing.append(mod)
            if missing:
                raise QgsProcessingException(
                    "Modules Python manquants : {}. Installez-les via la console "
                    "OSGeo4W ou le terminal Python de QGIS, par exemple :\n"
                    "python3 -m pip install {}".format(", ".join(missing), " ".join(missing))
                )

            def progress_cb(pct, msg):
                feedback.setProgress(pct)
                feedback.pushInfo(msg)

            feedback.pushInfo(f"Fichier KMZ : {kmz_path}")
            feedback.pushInfo(f"Tolérance : {tolerance} m")

            try:
                build_global_excel(kmz_path, out_path, tolerance_m=tolerance, progress_cb=progress_cb)
            except Exception as exc:
                raise QgsProcessingException(f"Échec de l'extraction : {exc}")

            feedback.pushInfo(f"Fichier généré : {out_path}")
            return {self.OUTPUT_XLSX: out_path}
